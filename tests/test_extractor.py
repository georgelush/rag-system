"""
tests/test_extractor.py — Unit tests for format extraction and MIME detection.
"""

import io
import json
import zipfile

import pytest

from rag.extractor import (
    ExtractionResult,
    _clean_html,
    _detect_mime,
    _extract_csv,
    _extract_json,
    _extract_xml,
    _extract_epub,
    extract,
)


# ── MIME detection ─────────────────────────────────────────────────────────────

def test_mime_hint_takes_priority():
    assert _detect_mime(b"", mime_hint="application/pdf") == "application/pdf"


def test_mime_extension_fallback_pdf():
    assert _detect_mime(b"", filename="report.pdf") == "application/pdf"


def test_mime_extension_fallback_docx():
    result = _detect_mime(b"", filename="document.docx")
    assert "wordprocessingml" in result


def test_mime_extension_fallback_csv():
    assert _detect_mime(b"", filename="data.csv") == "text/csv"


def test_mime_magic_bytes_pdf():
    assert _detect_mime(b"%PDF-1.4 ...") == "application/pdf"


def test_mime_magic_bytes_html():
    assert _detect_mime(b"<!DOCTYPE html><html>...") == "text/html"


def test_mime_plain_text_fallback():
    assert _detect_mime(b"hello world") == "text/plain"


def test_mime_audio_mp3():
    assert _detect_mime(b"", filename="audio.mp3") == "audio/mpeg"


# ── HTML extraction ────────────────────────────────────────────────────────────

def test_html_strips_tags():
    html = b"<html><body><p>Hello <b>world</b></p></body></html>"
    text = _clean_html(html)
    assert "Hello" in text
    assert "world" in text
    assert "<" not in text


def test_html_removes_script_tags():
    html = b"<html><script>alert('xss')</script><p>Safe content</p></html>"
    text = _clean_html(html)
    assert "alert" not in text
    assert "Safe content" in text


def test_html_removes_style_tags():
    html = b"<html><style>.red{color:red}</style><p>Visible</p></html>"
    text = _clean_html(html)
    assert ".red" not in text
    assert "Visible" in text


def test_html_decodes_entities():
    html = b"<p>AT&amp;T &lt;Corp&gt;</p>"
    text = _clean_html(html)
    assert "AT&T" in text
    assert "<Corp>" in text


# ── CSV extraction ─────────────────────────────────────────────────────────────

def test_csv_extracts_rows():
    csv_bytes = b"name,age\nAlice,30\nBob,25"
    text = _extract_csv(csv_bytes)
    assert "Alice" in text
    assert "Bob" in text
    assert "name" in text


def test_csv_skips_empty_rows():
    csv_bytes = b"col1,col2\n\nA,B\n\nC,D"
    text = _extract_csv(csv_bytes)
    assert "A" in text
    assert "C" in text


# ── JSON/JSONL extraction ──────────────────────────────────────────────────────

def test_json_object():
    data = b'{"title": "Document", "content": "Some text"}'
    text = _extract_json(data)
    assert "Document" in text
    assert "Some text" in text


def test_jsonl_multiple_objects():
    jsonl = b'{"id": 1, "text": "first line"}\n{"id": 2, "text": "second line"}'
    text = _extract_json(jsonl)
    assert "first line" in text
    assert "second line" in text


# ── XML extraction ─────────────────────────────────────────────────────────────

def test_xml_extracts_text_nodes():
    xml = b"<root><item>Alpha</item><item>Beta</item><desc>Gamma</desc></root>"
    text = _extract_xml(xml)
    assert "Alpha" in text
    assert "Beta" in text
    assert "Gamma" in text


def test_xml_invalid_falls_back_to_raw():
    bad_xml = b"not < valid > xml"
    text = _extract_xml(bad_xml)
    assert len(text) > 0  # doesn't raise, returns raw


# ── EPUB extraction ────────────────────────────────────────────────────────────

def _make_epub(*chapters: tuple[str, str]) -> bytes:
    """Create minimal EPUB bytes (ZIP of HTML files)."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for name, content in chapters:
            zf.writestr(name, f"<html><body><p>{content}</p></body></html>")
    return buf.getvalue()


def test_epub_extracts_html_chapters():
    epub = _make_epub(("ch1.html", "Chapter one text"), ("ch2.html", "Chapter two text"))
    text = _extract_epub(epub)
    assert "Chapter one text" in text
    assert "Chapter two text" in text


def test_epub_sorted_chapter_order():
    epub = _make_epub(("ch2.html", "Second"), ("ch1.html", "First"))
    text = _extract_epub(epub)
    # ch1 should appear before ch2 (alphabetical sort)
    assert text.index("First") < text.index("Second")


def test_epub_invalid_zip_returns_raw():
    text = _extract_epub(b"not a zip file")
    assert len(text) > 0


# ── async extract() dispatcher ─────────────────────────────────────────────────

async def test_extract_plain_text():
    result = await extract(b"Hello world", filename="file.txt")
    assert result.text == "Hello world"
    assert result.pages is None


async def test_extract_html():
    result = await extract(b"<p>Test content</p>", filename="page.html")
    assert "Test content" in result.text
    assert result.mime_type == "text/html"


async def test_extract_csv():
    result = await extract(b"col1,col2\nA,B\nC,D", filename="data.csv")
    assert "A" in result.text
    assert result.mime_type == "text/csv"


async def test_extract_json():
    result = await extract(b'{"key": "value"}', filename="data.json")
    assert "value" in result.text


async def test_extract_jsonl():
    data = b'{"a": 1}\n{"b": 2}'
    result = await extract(data, filename="data.jsonl")
    assert "1" in result.text


async def test_extract_xml():
    result = await extract(b"<root><node>Content</node></root>", filename="data.xml")
    assert "Content" in result.text


async def test_extract_epub_via_dispatcher():
    epub = _make_epub(("chapter.html", "EPUB chapter content"))
    result = await extract(epub, mime_hint="application/epub+zip")
    assert "EPUB chapter content" in result.text


async def test_extract_pdf_returns_pages():
    from unittest.mock import MagicMock, patch

    mock_page1 = MagicMock()
    mock_page1.extract_text.return_value = "Page 1 content"
    mock_page2 = MagicMock()
    mock_page2.extract_text.return_value = "Page 2 content"

    mock_reader = MagicMock()
    mock_reader.pages = [mock_page1, mock_page2]

    with patch("pypdf.PdfReader", return_value=mock_reader):
        result = await extract(b"%PDF-1.4 fake content", filename="doc.pdf")

    assert result.pages is not None
    assert len(result.pages) == 2
    assert "Page 1 content" in result.text
    assert "Page 2 content" in result.text
    assert result.mime_type == "application/pdf"


async def test_extract_unknown_extension_returns_text():
    result = await extract(b"raw bytes here", filename="file.unknown")
    assert "raw bytes" in result.text


async def test_extract_mime_hint_overrides_extension():
    result = await extract(b"<p>HTML</p>", filename="file.txt", mime_hint="text/html")
    assert "HTML" in result.text
    assert result.mime_type == "text/html"


# ── DOCX extraction ────────────────────────────────────────────────────────────

def _make_docx(paragraphs: list[str]) -> bytes:
    """Create a minimal .docx in-memory via python-docx."""
    import docx as _docx
    doc = _docx.Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_extracts_paragraphs():
    from rag.extractor import _extract_docx
    raw = _make_docx(["First paragraph.", "Second paragraph."])
    text = _extract_docx(raw)
    assert "First paragraph." in text
    assert "Second paragraph." in text


def test_docx_skips_empty_paragraphs():
    from rag.extractor import _extract_docx
    raw = _make_docx(["Non-empty", "", "Also non-empty"])
    text = _extract_docx(raw)
    assert "Non-empty" in text
    assert "Also non-empty" in text


async def test_extract_dispatches_docx():
    raw = _make_docx(["Docx content here."])
    result = await extract(raw, filename="report.docx")
    assert "Docx content here." in result.text
    assert "wordprocessingml" in result.mime_type


# ── PPTX extraction ────────────────────────────────────────────────────────────

def _make_pptx(slides: list[list[str]]) -> bytes:
    """Create a minimal .pptx in-memory via python-pptx."""
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    blank_layout = prs.slide_layouts[6]  # blank layout
    for slide_texts in slides:
        slide = prs.slides.add_slide(blank_layout)
        for i, text in enumerate(slide_texts):
            txBox = slide.shapes.add_textbox(Inches(i), Inches(0), Inches(3), Inches(1))
            txBox.text_frame.text = text
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_pptx_extracts_slide_text():
    from rag.extractor import _extract_pptx
    raw = _make_pptx([["Slide one title", "Slide one body"], ["Slide two"]])
    text = _extract_pptx(raw)
    assert "Slide one title" in text
    assert "Slide two" in text


def test_pptx_skips_empty_slides():
    from rag.extractor import _extract_pptx
    raw = _make_pptx([["Content slide"], []])
    text = _extract_pptx(raw)
    assert "Content slide" in text


async def test_extract_dispatches_pptx():
    raw = _make_pptx([["PowerPoint slide text."]])
    result = await extract(raw, filename="deck.pptx")
    assert "PowerPoint slide text." in result.text
    assert "presentationml" in result.mime_type


# ── XLSX extraction ────────────────────────────────────────────────────────────

def _make_xlsx(sheets: dict[str, list[list]]) -> bytes:
    """Create a minimal .xlsx in-memory via openpyxl."""
    import openpyxl
    wb = openpyxl.Workbook()
    first = True
    for sheet_name, rows in sheets.items():
        if first:
            ws = wb.active
            ws.title = sheet_name
            first = False
        else:
            ws = wb.create_sheet(sheet_name)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_xlsx_extracts_cell_values():
    from rag.extractor import _extract_xlsx
    raw = _make_xlsx({"Sheet1": [["Name", "Score"], ["Alice", 95], ["Bob", 87]]})
    text = _extract_xlsx(raw)
    assert "Alice" in text
    assert "95" in text
    assert "Bob" in text


def test_xlsx_includes_sheet_header():
    from rag.extractor import _extract_xlsx
    raw = _make_xlsx({"MySheet": [["Cell"]]})
    text = _extract_xlsx(raw)
    assert "MySheet" in text


def test_xlsx_skips_empty_rows():
    from rag.extractor import _extract_xlsx
    raw = _make_xlsx({"Sheet1": [["Value"], [], ["Another"]]})
    text = _extract_xlsx(raw)
    assert "Value" in text
    assert "Another" in text


async def test_extract_dispatches_xlsx():
    raw = _make_xlsx({"Sales": [["Q1", 100], ["Q2", 200]]})
    result = await extract(raw, filename="data.xlsx")
    assert "Q1" in result.text
    assert "spreadsheetml" in result.mime_type


# ── Audio / Whisper ────────────────────────────────────────────────────────────

async def test_extract_audio_calls_whisper_and_returns_transcript():
    """Audio bytes → Whisper API called → transcript returned as text."""
    from unittest.mock import AsyncMock, MagicMock, patch

    mock_transcript = MagicMock()
    mock_transcript.text = "This is the transcribed audio content."

    with patch("rag.extractor.openai_client") as mock_client:
        mock_client.audio.transcriptions.create = AsyncMock(return_value=mock_transcript)
        result = await extract(b"\xff\xfb fake mp3 bytes", filename="recording.mp3")

    assert result.text == "This is the transcribed audio content."
    assert result.mime_type == "audio/mpeg"
    mock_client.audio.transcriptions.create.assert_awaited_once()


async def test_extract_audio_passes_correct_model_and_filename():
    """Whisper call uses configured WHISPER_MODEL and original filename."""
    from unittest.mock import AsyncMock, MagicMock, patch
    import rag.extractor as _extractor

    mock_transcript = MagicMock()
    mock_transcript.text = "Transcript."
    call_kwargs: dict = {}

    async def _capture(**kwargs):
        call_kwargs.update(kwargs)
        return mock_transcript

    with patch.object(_extractor, "openai_client") as mock_client:
        mock_client.audio.transcriptions.create = _capture
        await extract(b"audio bytes", filename="lecture.wav")

    assert call_kwargs["model"] == _extractor.WHISPER_MODEL
    # file tuple: (filename, bytes, mime)
    assert call_kwargs["file"][0] == "lecture.wav"
    assert call_kwargs["file"][2] == "audio/wav"


async def test_extract_audio_via_mime_hint():
    """mime_hint=audio/mpeg triggers Whisper even with non-audio extension."""
    from unittest.mock import AsyncMock, MagicMock, patch

    mock_transcript = MagicMock()
    mock_transcript.text = "Audio via hint."

    with patch("rag.extractor.openai_client") as mock_client:
        mock_client.audio.transcriptions.create = AsyncMock(return_value=mock_transcript)
        result = await extract(b"bytes", filename="file.bin", mime_hint="audio/mpeg")

    assert "Audio via hint." in result.text


async def test_extract_wav_dispatches_to_whisper():
    """WAV extension → Whisper path, not text fallback."""
    from unittest.mock import AsyncMock, MagicMock, patch

    mock_transcript = MagicMock()
    mock_transcript.text = "WAV transcript."

    with patch("rag.extractor.openai_client") as mock_client:
        mock_client.audio.transcriptions.create = AsyncMock(return_value=mock_transcript)
        result = await extract(b"RIFF fake wav", filename="clip.wav")

    assert result.text == "WAV transcript."
