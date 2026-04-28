"""
rag/extractor.py — Extract plain text from all supported file formats.

Supported:
  Documents : PDF, HTML, TXT, MD, DOCX, PPTX, EPUB
  Data      : CSV, JSON, JSONL, XLSX, XML
  Audio     : MP3, MP4, WAV, OGG, FLAC, WEBM, M4A  (via Whisper → LLM_PROXY)

Returns an ExtractionResult with:
  text      — full plain text (all formats)
  pages     — per-page strings (PDF only, used by PageChunker)
  mime_type — detected MIME type string
"""

import csv
import io
import json
import os
import re
import xml.etree.ElementTree as ET
import zipfile
from dataclasses import dataclass, field

from rag.openai_client import openai_client

WHISPER_MODEL = os.environ.get("WHISPER_MODEL", "whisper-1")

# ── MIME / Extension maps ──────────────────────────────────────────────────────

_EXT_TO_MIME: dict[str, str] = {
    ".pdf":   "application/pdf",
    ".html":  "text/html",
    ".htm":   "text/html",
    ".txt":   "text/plain",
    ".md":    "text/markdown",
    ".docx":  "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pptx":  "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".xlsx":  "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".csv":   "text/csv",
    ".json":  "application/json",
    ".jsonl": "application/jsonlines",
    ".xml":   "application/xml",
    ".epub":  "application/epub+zip",
    ".mp3":   "audio/mpeg",
    ".mp4":   "audio/mp4",
    ".wav":   "audio/wav",
    ".ogg":   "audio/ogg",
    ".flac":  "audio/flac",
    ".webm":  "video/webm",
    ".m4a":   "audio/m4a",
}

_AUDIO_MIMES = {
    "audio/mpeg", "audio/mp4", "audio/wav",
    "audio/ogg", "audio/flac", "video/webm", "audio/m4a",
}

_AUDIO_EXTS = {
    ".mp3", ".mp4", ".wav", ".ogg", ".flac", ".webm", ".m4a",
}

# ── Result type ────────────────────────────────────────────────────────────────

@dataclass
class ExtractionResult:
    text:      str
    pages:     list[str] | None = field(default=None)  # PDF only
    mime_type: str = ""


# ── MIME detection ─────────────────────────────────────────────────────────────

def _detect_mime(raw: bytes, filename: str = "", mime_hint: str = "") -> str:
    """Detect MIME: hint → file extension → magic bytes → plain-text fallback."""
    if mime_hint:
        return mime_hint.lower().split(";")[0].strip()

    ext = os.path.splitext(filename)[1].lower() if filename else ""
    if ext in _EXT_TO_MIME:
        return _EXT_TO_MIME[ext]

    # Magic bytes
    if raw[:4] == b"%PDF":
        return "application/pdf"
    if raw[:2] == b"PK":
        # ZIP-based (DOCX / XLSX / PPTX / EPUB) — can't distinguish without extension
        return "application/zip"
    if raw.lstrip()[:5].lower() in (b"<!doc", b"<html"):
        return "text/html"

    return "text/plain"


# ── Format extractors ──────────────────────────────────────────────────────────

def _clean_html(raw: bytes) -> str:
    text = raw.decode("utf-8", errors="replace")
    text = re.sub(r"<script[^>]*>.*?</script>", " ", text, flags=re.DOTALL | re.IGNORECASE)
    text = re.sub(r"<style[^>]*>.*?</style>",  " ", text, flags=re.DOTALL | re.IGNORECASE)
    text = re.sub(r"<[^>]+>", " ", text)
    text = (text
            .replace("&nbsp;", " ")
            .replace("&amp;",  "&")
            .replace("&lt;",   "<")
            .replace("&gt;",   ">")
            .replace("&quot;", '"'))
    return re.sub(r"\s+", " ", text).strip()


def _extract_pdf(raw: bytes) -> ExtractionResult:
    import pypdf
    reader = pypdf.PdfReader(io.BytesIO(raw))
    pages = [page.extract_text() or "" for page in reader.pages]
    return ExtractionResult(
        text="\n\n".join(p for p in pages if p.strip()),
        pages=pages,
        mime_type="application/pdf",
    )


def _extract_docx(raw: bytes) -> str:
    import docx
    doc = docx.Document(io.BytesIO(raw))
    return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _extract_pptx(raw: bytes) -> str:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(raw))
    slides = []
    for slide in prs.slides:
        texts = [
            shape.text.strip()
            for shape in slide.shapes
            if hasattr(shape, "text") and shape.text.strip()
        ]
        if texts:
            slides.append("\n".join(texts))
    return "\n\n".join(slides)


def _extract_xlsx(raw: bytes) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(c.strip() for c in cells):
                rows.append(", ".join(cells))
        if rows:
            sheets.append(f"[Sheet: {ws.title}]\n" + "\n".join(rows))
    return "\n\n".join(sheets)


def _extract_csv(raw: bytes) -> str:
    text = raw.decode("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = [", ".join(row) for row in reader if any(c.strip() for c in row)]
    return "\n".join(rows)


def _extract_json(raw: bytes) -> str:
    text = raw.decode("utf-8", errors="replace").strip()
    # JSONL: one JSON object per line, no outer array bracket
    if "\n" in text and not text.startswith("["):
        lines = []
        for line in text.splitlines():
            line = line.strip()
            if not line:
                continue
            try:
                lines.append(json.dumps(json.loads(line), ensure_ascii=False))
            except json.JSONDecodeError:
                lines.append(line)
        return "\n".join(lines)
    # Regular JSON
    try:
        return json.dumps(json.loads(text), ensure_ascii=False, indent=2)
    except json.JSONDecodeError:
        return text


def _extract_xml(raw: bytes) -> str:
    try:
        root = ET.fromstring(raw)
        parts = [e.text.strip() for e in root.iter() if e.text and e.text.strip()]
        return "\n".join(parts)
    except ET.ParseError:
        return raw.decode("utf-8", errors="replace")


def _extract_epub(raw: bytes) -> str:
    """Extract text from EPUB (ZIP of HTML files) — no external dependencies."""
    texts: list[str] = []
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as zf:
            html_files = sorted(
                n for n in zf.namelist()
                if n.endswith((".html", ".htm", ".xhtml"))
                and not n.startswith("__MACOSX")
            )
            for name in html_files:
                extracted = _clean_html(zf.read(name))
                if extracted:
                    texts.append(extracted)
    except zipfile.BadZipFile:
        return raw.decode("utf-8", errors="replace")
    return "\n\n".join(texts)


async def _extract_audio(raw: bytes, filename: str, mime: str) -> str:
    """Transcribe audio/video via Whisper through LLM_PROXY."""
    transcript = await openai_client.audio.transcriptions.create(
        model=WHISPER_MODEL,
        file=(filename or "audio.mp3", raw, mime),
    )
    return transcript.text


# ── Main entry point ───────────────────────────────────────────────────────────

async def extract(
    raw: bytes,
    filename: str = "",
    mime_hint: str = "",
) -> ExtractionResult:
    """Detect file format and extract plain text.

    Args:
        raw:       Raw file bytes (from URL fetch or multipart upload).
        filename:  Original filename — drives extension-based MIME detection.
        mime_hint: MIME type hint from IngestRequest.mime_type_hint.

    Returns:
        ExtractionResult with .text (always) and .pages (PDF only).

    Raises:
        ValueError: propagated from format parsers on corrupt/unsupported files.
    """
    mime = _detect_mime(raw, filename, mime_hint)
    ext  = os.path.splitext(filename)[1].lower() if filename else ""

    # Audio → Whisper transcription
    if mime in _AUDIO_MIMES or ext in _AUDIO_EXTS:
        text = await _extract_audio(raw, filename or "audio.mp3", mime)
        return ExtractionResult(text=text, mime_type=mime)

    # PDF — special: returns per-page list for PageChunker
    if mime == "application/pdf" or raw[:4] == b"%PDF":
        return _extract_pdf(raw)

    # HTML
    if "html" in mime:
        return ExtractionResult(text=_clean_html(raw), mime_type=mime)

    # DOCX
    if "wordprocessingml" in mime or ext == ".docx":
        return ExtractionResult(text=_extract_docx(raw), mime_type=mime)

    # PPTX
    if "presentationml" in mime or ext == ".pptx":
        return ExtractionResult(text=_extract_pptx(raw), mime_type=mime)

    # XLSX
    if "spreadsheetml" in mime or ext == ".xlsx":
        return ExtractionResult(text=_extract_xlsx(raw), mime_type=mime)

    # CSV
    if "csv" in mime or ext == ".csv":
        return ExtractionResult(text=_extract_csv(raw), mime_type=mime)

    # JSON / JSONL
    if "json" in mime or ext in (".json", ".jsonl"):
        return ExtractionResult(text=_extract_json(raw), mime_type=mime)

    # XML
    if "xml" in mime or ext == ".xml":
        return ExtractionResult(text=_extract_xml(raw), mime_type=mime)

    # EPUB
    if "epub" in mime or ext == ".epub":
        return ExtractionResult(text=_extract_epub(raw), mime_type=mime)

    # Unknown ZIP-based — try EPUB structure, then plain text
    if mime == "application/zip" or raw[:2] == b"PK":
        try:
            text = _extract_epub(raw)
            if text.strip():
                return ExtractionResult(text=text, mime_type="application/epub+zip")
        except Exception:
            pass

    # Markdown / TXT / ultimate fallback
    return ExtractionResult(
        text=raw.decode("utf-8", errors="replace"),
        mime_type=mime,
    )
